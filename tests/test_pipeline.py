"""Pruebas de integridad y fallos; dobles de API, sin medir calidad del LLM."""
import copy
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openai import APIConnectionError
from openpyxl import load_workbook

from aqr.budget import Budget, BudgetStop, cost_micro
from aqr.data import load_cases
from aqr.export import export_workbook
from aqr.provider import OpenAIProvider, StageFailure
from aqr.review import import_review
from aqr.schemas import MODELS, Tipologia
from aqr.storage import atomic_json, project_lock, read_json
from aqr.validation import validate_stage
from aqr.workflow import build_graph, process_case

ROOT = Path(__file__).resolve().parents[1]
CONFIG = read_json(ROOT / "config/openai.json")
SOURCE = {"id": "radicado:2:nota", "archivo": "casos.xlsx", "hoja": "Datos", "fila": 2,
          "campo": "nota", "texto": "Solicito el extracto de mi libranza.", "ausente": False}
CITE = {"fuente_id": SOURCE["id"], "texto": SOURCE["texto"]}
CASE = {"radicado": "-123", "development_id": "D01", "sources": [SOURCE], "customer_candidates": [{}],
        "join": {"candidate_count": 1, "ambiguous": False, "conflicts": {}}, "source_fingerprints": []}
OUTPUTS = {
    "clasificacion": {"tipologia": Tipologia.EXTRACTO.value, "causalidad": "Solicitud de extracto",
        "solicitud_principal": "Obtener extracto", "solicitudes_secundarias": [], "evidencia": [CITE],
        "justificacion": "La nota solicita el extracto de libranza.", "asignacion_forzada": False,
        "requiere_revision": False, "razones_revision": []},
    "investigacion": {"hallazgos": [{"id": "H1", "tipo": "declaracion_cliente",
        "descripcion": "El cliente solicita extracto de libranza.", "evidencia": [CITE]}],
        "causa_documentada_ids": [], "datos_pendientes": ["Extracto solicitado no disponible"],
        "limitaciones": [], "requiere_revision": True, "razones_revision": ["Falta extracto"]},
    "respuesta": {"parrafos": [
        {"texto": "Según su solicitud, requiere el extracto de libranza.", "tipo": "hecho", "evidencia_ids": ["H1"], "pendiente_indices": []},
        {"texto": "El extracto solicitado no está disponible en la información revisada.", "tipo": "pendiente", "evidencia_ids": [], "pendiente_indices": [0]}],
        "requiere_revision": True, "razones_revision": ["Falta extracto"]},
}


class FakeProvider:
    def __init__(self, outputs=None, stop=None):
        self.outputs = copy.deepcopy(outputs or OUTPUTS)
        self.calls = []
        self.stop = stop

    def generate(self, stage, *args):
        self.calls.append(stage)
        if self.stop == stage:
            raise BudgetStop("Presupuesto agotado de prueba")
        return json.dumps(self.outputs[stage]), {"usage": {"input_tokens": 10, "output_tokens": 10}}


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_three_separate_stages_and_cached_resume(self):
        provider = FakeProvider()
        graph = build_graph(provider, self.root / "stages", "v1")
        first = process_case(graph, CASE)
        second = process_case(graph, CASE)
        self.assertEqual(provider.calls, list(MODELS))
        self.assertEqual(first["status"], "completado")
        self.assertEqual(first["draft"], second["draft"])
        self.assertTrue(first["requires_review"])
        self.assertTrue(all(x["cached"] for x in second["metadata"].values()))

    def test_version_change_invalidates_cache(self):
        provider = FakeProvider()
        process_case(build_graph(provider, self.root / "stages", "v1"), CASE)
        process_case(build_graph(provider, self.root / "stages", "v2"), CASE)
        self.assertEqual(len(provider.calls), 6)

    def test_invented_citation_stops_before_investigation(self):
        outputs = copy.deepcopy(OUTPUTS)
        outputs["clasificacion"]["evidencia"][0]["texto"] = "Se devolvió el dinero"
        provider = FakeProvider(outputs)
        result = process_case(build_graph(provider, self.root, "v1"), CASE)
        self.assertEqual(provider.calls, ["clasificacion"])
        self.assertEqual(result["status"], "incompleto")
        self.assertIsNone(result["draft"])

    def test_budget_stop_preserves_previous_stage_and_resumes(self):
        provider = FakeProvider(stop="investigacion")
        result = process_case(build_graph(provider, self.root, "v1"), CASE)
        self.assertIn("clasificacion", result["outputs"])
        self.assertTrue(result["stop_batch"])
        self.assertEqual(result["status"], "interrumpido")
        fresh = FakeProvider()
        resumed = process_case(build_graph(fresh, self.root, "v1"), CASE)
        self.assertEqual(fresh.calls, ["investigacion", "respuesta"])
        self.assertEqual(resumed["status"], "completado")

    def test_missing_fact_reference_and_pending_omission(self):
        answer = copy.deepcopy(OUTPUTS["respuesta"])
        answer["parrafos"] = [answer["parrafos"][0]]
        answer["parrafos"][0]["evidencia_ids"] = ["inventado"]
        errors = validate_stage("respuesta", answer, CASE, OUTPUTS)
        self.assertGreaterEqual(len(errors), 2)

    def test_cause_cannot_reference_customer_allegation(self):
        investigation = copy.deepcopy(OUTPUTS["investigacion"])
        investigation["causa_documentada_ids"] = ["H1"]
        self.assertTrue(validate_stage("investigacion", investigation, CASE, OUTPUTS))

    def test_budget_unknown_charge_survives_restart(self):
        budget = Budget(self.root / "budget.json", CONFIG)
        budget.reserve("one", 1000, 1000, {})
        budget.uncertain("one", "Timeout")
        restarted = Budget(self.root / "budget.json", CONFIG)
        self.assertEqual(restarted.total(), cost_micro(1000, 1000, CONFIG))
        self.assertGreater(restarted.summary()["reserved_or_uncertain_usd"], 0)

    def test_budget_cap_prevents_call_and_settlement_releases_excess(self):
        config = {**CONFIG, "budget_usd": "0.002"}
        budget = Budget(self.root / "budget.json", config)
        budget.reserve("one", 1000, 1000, {})
        with self.assertRaises(BudgetStop):
            budget.reserve("two", 1000, 1000, {})
        budget.settle("one", {"input_tokens": 100, "output_tokens": 100}, "response")
        budget.reserve("two", 1000, 1000, {})
        self.assertLessEqual(budget.total(), 2000)

    def test_cannot_increase_authorized_cap_or_reset_missing_ledger(self):
        with self.assertRaises(ValueError):
            Budget(self.root / "budget.json", {**CONFIG, "budget_usd": "4"})
        atomic_json(self.root / "stages/classification/x.json", {})
        with self.assertRaises(BudgetStop):
            Budget(self.root / "budget.json", CONFIG)

    def test_second_process_cannot_take_lock(self):
        with project_lock(self.root / ".lock"):
            with self.assertRaises(RuntimeError):
                with project_lock(self.root / ".lock"):
                    pass

    def test_real_data_includes_hidden_rows_and_keeps_ambiguous_customers(self):
        cases = load_cases(ROOT)
        self.assertEqual(len(cases), 275)
        self.assertEqual(sum(bool(c["development_id"]) for c in cases), 20)
        self.assertEqual(sum(len(c["customer_candidates"]) for c in cases), 345)
        self.assertTrue(all(isinstance(c["radicado"], str) for c in cases))
        row19 = next(c for c in cases if c["radicado_source"]["source_row"] == 19)
        self.assertEqual(len(row19["join"]["conflicts"]["numero_de_credito"]), 2)

    def test_export_preserves_formula_like_text_and_blank_human_fields(self):
        provider = FakeProvider()
        result = process_case(build_graph(provider, self.root / "stages", "v1"), CASE)
        result["outputs"]["clasificacion"]["causalidad"] = '=HYPERLINK("https://invalid.test")'
        run = {"run_id": "SYNTHETIC_TEST_ONLY", "workflow_fingerprint": "test", "results": [result],
               "budget": {}, "source_fingerprints": []}
        output = self.root / "test.xlsx"
        export_workbook(run, [CASE], output)
        wb = load_workbook(output)
        try:
            self.assertEqual(wb["Resultados"]["E2"].data_type, "s")
            self.assertTrue(all(wb["Revision"].cell(2, c).value is None for c in range(5, 12)))
        finally:
            wb.close()
        with self.assertRaises(FileExistsError):
            export_workbook(run, [CASE], output)

    def test_sdk_request_caps_no_tools_and_receipt_recovery(self):
        budget = Budget(self.root / "budget.json", CONFIG)
        response = Mock()
        response.output_text = json.dumps(OUTPUTS["clasificacion"])
        response.model_dump.return_value = {"id": "r1", "status": "completed", "model": CONFIG["model"],
            "usage": {"input_tokens": 100, "output_tokens": 120}}
        client = Mock()
        client.responses.input_tokens.count.return_value = SimpleNamespace(input_tokens=100)
        client.responses.create.return_value = response
        provider = OpenAIProvider(CONFIG, budget, self.root / "receipts", client=client)
        args = ("clasificacion", "request1", "Instructions", {"fuentes": [SOURCE]}, MODELS["clasificacion"].model_json_schema())
        first = provider.generate(*args)
        again = provider.generate(*args)
        self.assertEqual(client.responses.create.call_count, 1)
        self.assertEqual(first, again)
        params = client.responses.create.call_args.kwargs
        self.assertFalse(params["store"])
        self.assertNotIn("tools", params)
        self.assertTrue(params["text"]["format"]["strict"])
        self.assertEqual(params["max_output_tokens"], 1600)
        self.assertEqual(budget.data["calls"]["request1-0"]["state"], "settled")

    def test_uncertain_previous_call_is_not_repeated(self):
        budget = Budget(self.root / "budget.json", CONFIG)
        budget.reserve("key-0", 500, 1600, {})
        provider = OpenAIProvider(CONFIG, budget, self.root / "receipts", client=Mock())
        with self.assertRaises(StageFailure):
            provider.generate("clasificacion", "key", "", {}, {})
        provider.client.responses.create.assert_not_called()

    def test_import_requires_explicit_acceptance(self):
        result = process_case(build_graph(FakeProvider(), self.root / "stages", "v1"), CASE)
        run = {"run_id": "TEST_ONLY", "scope": "development", "workflow_fingerprint": "test", "results": [result],
               "budget": {}, "source_fingerprints": []}
        workbook = self.root / "review.xlsx"
        export_workbook(run, [CASE], workbook)
        blank = import_review(workbook, run, self.root / "blank.json")
        self.assertEqual(blank["reviewed_count"], 0)
        self.assertIsNone(blank["annotations"][0]["human_label"])
        wb = load_workbook(workbook)
        wb["Revision"]["E2"] = "Aceptar"
        wb["Revision"]["C2"] = "Modified suggestion must not become the accepted label"
        wb.save(workbook)
        wb.close()
        accepted = import_review(workbook, run, self.root / "accepted.json")
        self.assertEqual(accepted["reviewed_count"], 1)
        self.assertEqual(accepted["annotations"][0]["human_label"]["tipologia"], Tipologia.EXTRACTO.value)

    def test_recover_second_attempt_receipt_without_repeating_first(self):
        budget = Budget(self.root / "budget.json", CONFIG)
        budget.reserve("key-0", 500, 1600, {})
        budget.uncertain("key-0", "Timeout")
        budget.reserve("key-1", 500, 1600, {})
        atomic_json(self.root / "receipts/key-1.json", {"output_text": "{}", "latency_seconds": 1,
            "response": {"id": "second", "status": "completed", "usage": {"input_tokens": 100, "output_tokens": 100}}})
        provider = OpenAIProvider(CONFIG, budget, self.root / "receipts", client=Mock())
        text, _ = provider.generate("clasificacion", "key", "", {}, {})
        self.assertEqual(text, "{}")
        provider.client.responses.create.assert_not_called()
        self.assertEqual(budget.data["calls"]["key-0"]["state"], "uncertain")


if __name__ == "__main__":
    unittest.main()
