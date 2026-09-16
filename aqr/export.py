"""Excel de resultados y revisión asistida; no sobreescribe etiquetas humanas."""
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

from scripts.crear_excel_revision import chunks, long_height, text_cell
from .schemas import Tipologia


def sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "C2"
    for i, (header, width) in enumerate(zip(headers, widths, strict=True), 1):
        text_cell(ws.cell(1, i), header, color="FFFFFF", bold=True)
        ws.cell(1, i).fill = PatternFill("solid", fgColor="243C5A")
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    ws.row_dimensions[1].height = 42
    return ws


def row(ws, values, height=65):
    index = ws.max_row + 1
    for i, value in enumerate(values, 1):
        if value is not None:
            text_cell(ws.cell(index, i), value)
        else:
            ws.cell(index, i).alignment = Alignment(vertical="top", wrap_text=True)
        if index % 2 == 0:
            ws.cell(index, i).fill = PatternFill("solid", fgColor="EDF2F7")
    ws.row_dimensions[index].height = height
    return index


def export_workbook(run, cases, output: Path):
    if output.exists():
        raise FileExistsError("El Excel existe; usa una ruta nueva para conservar las revisiones")
    if any(c["source_fingerprints"] != run["source_fingerprints"] for c in cases):
        raise ValueError("La ejecución y los expedientes no corresponden a las mismas fuentes")
    case_map = {c["radicado"]: c for c in cases}
    results = run["results"]
    if not results or len({r["radicado"] for r in results}) != len(results):
        raise ValueError("Resultados vacíos o radicados repetidos")
    wb = Workbook()
    wb.remove(wb.active)
    ws = sheet(wb, "Resultados", ["Caso", "Radicado", "Estado técnico", "Tipología propuesta",
        "Causalidad propuesta", "Asignación forzada", "Revisión requerida", "Detalle"],
        [10, 18, 20, 42, 54, 18, 20, 22])
    review = sheet(wb, "Revision", ["Caso", "Radicado", "Tipología propuesta", "Causalidad propuesta",
        "Decisión humana", "Tipología corregida", "Causalidad corregida", "Evidencia / justificación humana",
        "Investigación revisada", "Carta revisada", "Observaciones"], [10, 18, 42, 52, 22, 42, 52, 58, 22, 22, 58])
    detail = sheet(wb, "Detalle", ["Caso / sección", "Contenido"], [32, 124])
    evidence = sheet(wb, "Evidencias", ["Caso", "Hallazgo / etapa", "Fuente", "Ubicación", "Cita literal"],
                     [10, 27, 43, 48, 110])
    notes = sheet(wb, "Expedientes", ["Caso / campo", "Contenido original"], [32, 124])
    meta = sheet(wb, "Metadatos", ["Campo", "Valor"], [35, 110])
    instructions = [
        ("Uso", "Revisión asistida de desarrollo. Las sugerencias no son etiquetas humanas."),
        ("Decisión humana", "Aceptar o Corregir, explícitamente. Sin decisión, el caso queda sin revisar."),
        ("Correcciones", "Si corriges, completa tipología y causalidad finales, incluso si una coincide con la propuesta."),
        ("Investigación y carta", "Marcar Correcta o Requiere cambios; documentar correcciones en Observaciones."),
        ("Causalidad", "Motivo específico sustentado en la nota, no causa raíz demostrada ni catálogo oficial."),
        ("Limitación", "Citas verificadas automáticamente no garantizan interpretación correcta. Todos los textos son borradores."),
        ("Run ID", run["run_id"]), ("Versión del flujo", run["workflow_fingerprint"]),
        ("Presupuesto", json.dumps(run["budget"], ensure_ascii=False)),
    ]
    for name, value in instructions:
        row(meta, [name, value])
    for i, tip in enumerate(Tipologia, 15):
        text_cell(meta.cell(i, 1), "Tipología permitida")
        text_cell(meta.cell(i, 2), tip.value)
    for column, formula in [("E", '"Aceptar,Corregir"'), ("F", "'Metadatos'!$B$15:$B$20"),
                            ("I", '"Correcta,Requiere cambios"'), ("J", '"Correcta,Requiere cambios"')]:
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        validation.errorTitle = "Valor no permitido"
        validation.error = "Selecciona una opción de la lista."
        validation.showErrorMessage = True
        review.add_data_validation(validation)
        validation.add(f"{column}2:{column}{len(results) + 1}")
    for result in results:
        case = case_map[result["radicado"]]
        cid = result["development_id"] or result["radicado"]
        outputs = result["outputs"]
        classification = outputs.get("clasificacion", {})
        target = detail.max_row + 1
        r = row(ws, [cid, result["radicado"], result["status"], classification.get("tipologia"),
            classification.get("causalidad"), "Sí" if classification.get("asignacion_forzada") else "No" if classification else "Pendiente",
            "Sí" if result["requires_review"] else "No", "Ver detalle"])
        ws.cell(r, 8).hyperlink = Hyperlink(ref=ws.cell(r, 8).coordinate, location=f"'Detalle'!A{target}")
        rr = row(review, [cid, result["radicado"], classification.get("tipologia"),
                          classification.get("causalidad"), *([None] * 7)], 90)
        for column in range(5, 12):
            review.cell(rr, column).fill = PatternFill("solid", fgColor="FFF2CC")
        row(detail, [cid, f"Radicado {result['radicado']} — {result['status']}"], 30)
        sections = [("Clasificación", json.dumps(classification, ensure_ascii=False, indent=2)),
                    ("Cruce de cliente", json.dumps(case["join"], ensure_ascii=False, indent=2)),
                    ("Investigación", json.dumps(outputs.get("investigacion", {}), ensure_ascii=False, indent=2)),
                    ("Borrador", result["draft"] or "No generado; consultar estado técnico."),
                    ("Revisión / errores", "\n".join(result["review_reasons"]))]
        for name, content in sections:
            for part in chunks(content):
                row(detail, [name, part], long_height(part))
        source_map = {s["id"]: s for s in case["sources"]}
        cited = [("Clasificación", c) for c in classification.get("evidencia", [])]
        for finding in outputs.get("investigacion", {}).get("hallazgos", []):
            cited.extend((finding["id"], c) for c in finding["evidencia"])
        for label, cite in cited:
            source = source_map[cite["fuente_id"]]
            for part in chunks(cite["texto"]):
                row(evidence, [cid, label, source["archivo"],
                    f"{source['hoja']} / fila {source['fila']} / {source['campo']}", part], long_height(part, 85))
        note_row = notes.max_row + 1
        for source in case["sources"]:
            for part in chunks(source["texto"]):
                row(notes, [f"{cid} / {source['id']}", part], long_height(part))
        detail.cell(target, 1).hyperlink = Hyperlink(ref=detail.cell(target, 1).coordinate,
                                                    location=f"'Expedientes'!A{note_row}")
    for sh in (ws, review, evidence):
        sh.auto_filter.ref = sh.dimensions
    output.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive file creation prevents overwriting an annotation workbook even on a race.
    with output.open("xb") as stream:
        wb.save(stream)
    checked = load_workbook(output)
    try:
        if any(cell.data_type == "f" for sh in checked for cells in sh for cell in cells):
            raise ValueError("No se permiten fórmulas procedentes de datos o modelos")
        if any(checked["Revision"].cell(r, c).value is not None
               for r in range(2, len(results) + 2) for c in range(5, 12)):
            raise ValueError("La revisión humana debe comenzar vacía")
    finally:
        checked.close()
