"""Importación explícita de etiquetas humanas; jamás asumir aceptación por omisión."""
from pathlib import Path

from openpyxl import load_workbook

from .schemas import Tipologia
from .storage import atomic_json


def import_review(workbook: Path, run: dict, output: Path):
    if output.exists():
        raise FileExistsError("Las revisiones existentes no se sobreescriben")
    if run["scope"] != "development":
        raise ValueError("Este importador es solo para revisión asistida de desarrollo")
    wb = load_workbook(workbook, data_only=False)
    try:
        metadata = {r[0].value: r[1].value for r in wb["Metadatos"].iter_rows(min_row=2)}
        if metadata.get("Run ID") != run["run_id"] or metadata.get("Versión del flujo") != run["workflow_fingerprint"]:
            raise ValueError("El Excel pertenece a otra ejecución o versión")
        expected = {r["radicado"]: r for r in run["results"]}
        annotations, seen = [], set()
        for cells in wb["Revision"].iter_rows(min_row=2, max_col=11):
            values = [c.value for c in cells]
            if all(v is None for v in values):
                continue
            if any(c.data_type == "f" for c in cells):
                raise ValueError("La revisión debe contener valores, no fórmulas")
            radicado = str(values[1])
            if radicado not in expected or radicado in seen:
                raise ValueError("Radicado inesperado o repetido en revisión")
            seen.add(radicado)
            item = expected[radicado]
            if str(values[0]) != str(item["development_id"] or radicado):
                raise ValueError("El código de caso y el radicado no corresponden")
            decision = values[4]
            if decision is None:
                annotations.append({"radicado": radicado, "status": "sin_revisar", "human_label": None})
                continue
            if decision not in {"Aceptar", "Corregir"}:
                raise ValueError(f"Decisión no válida en {radicado}")
            original = item["outputs"].get("clasificacion")
            if decision == "Aceptar":
                if not original:
                    raise ValueError("No se puede aceptar una clasificación que no fue generada")
                # Usar resultado registrado, no columnas de sugerencias editables del Excel.
                tipologia, causalidad = original["tipologia"], original["causalidad"]
                if values[5] is not None or values[6] is not None:
                    raise ValueError("Hay correcciones pero la decisión dice Aceptar")
            else:
                tipologia, causalidad = values[5], values[6]
                if tipologia not in {t.value for t in Tipologia} or not str(causalidad or "").strip():
                    raise ValueError("Corregir requiere tipología y causalidad completas")
                if not str(values[7] or "").strip():
                    raise ValueError("Una corrección requiere evidencia o justificación humana")
            for value in values[8:10]:
                if value not in {None, "Correcta", "Requiere cambios"}:
                    raise ValueError("Resultado de revisión de investigación/carta no válido")
            annotations.append({"radicado": radicado, "status": "revisado", "decision": decision,
                "human_label": {"tipologia": tipologia, "causalidad": causalidad},
                "evidencia_humana": values[7], "revision_investigacion": values[8],
                "revision_carta": values[9], "observaciones": values[10]})
        if seen != set(expected):
            raise ValueError("Faltan casos de la ejecución en el Excel")
    finally:
        wb.close()
    result = {"schema_version": 1, "run_id": run["run_id"], "workflow_fingerprint": run["workflow_fingerprint"],
              "review_method": "assisted_development", "independent_test": False,
              "reviewed_count": sum(a["status"] == "revisado" for a in annotations), "annotations": annotations}
    atomic_json(output, result)
    return result
