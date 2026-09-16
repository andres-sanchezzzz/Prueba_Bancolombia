"""Build a blank human-review workbook from the selected development cases.

Uses openpyxl as explicitly selected by the user. Makes no model requests.
"""
from __future__ import annotations

import argparse
import json
import textwrap
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink


ROOT = Path(__file__).resolve().parents[1]
TIPOLOGIAS = [
    "Aclaraciones Cartera Hipotecaria",
    "Retiros por Pin Pad Cuenta Corriente",
    "Gravamen a Movimiento AFC",
    "Requerimiento de Extractos Libranza",
    "Doblemente Radicado",
    "No tramitado por falta de información",
]
NAVY = "243C5A"
AMBER = "FFF2CC"
PALE = "EDF2F7"
INK = "202A35"
GRAY = "606B78"


def text_cell(cell, value, *, color=INK, bold=False, size=11):
    """Write untrusted source text literally, including leading =, +, - or @."""
    cell.value = "(vacío)" if value is None else str(value)
    cell.data_type = "s"
    cell.font = Font(name="Arial", size=size, color=color, bold=bold)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def chunks(text: str, size: int = 600) -> list[str]:
    # No character is dropped. Break near whitespace when possible.
    parts = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            space = max(text.rfind(" ", start + size // 2, end), text.rfind("\n", start + size // 2, end))
            if space > start:
                end = space + 1
        parts.append(text[start:end])
        start = end
    return parts or [""]


def long_height(text: str, width: int = 100) -> float:
    lines = sum(max(1, len(textwrap.wrap(line, width=width, replace_whitespace=False)))
                for line in text.split("\n"))
    return min(409, max(30, lines * 15 + 14))


def section(sheet, row, title):
    for col in (1, 2):
        sheet.cell(row, col).fill = PatternFill("solid", fgColor=NAVY)
        sheet.cell(row, col).font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    sheet.cell(row, 1).value = title
    sheet.row_dimensions[row].height = 26


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=ROOT / "revision" / "desarrollo.json")
    parser.add_argument("--output", type=Path,
                        default=ROOT / "outputs" / "revision_inicial_20260915" / "Revision_20_casos.xlsx")
    args = parser.parse_args()
    if args.output.exists():
        raise FileExistsError("Refusing to overwrite a workbook that may contain human annotations")
    payload = json.loads(args.input.read_text(encoding="utf-8"))
    cases = payload["cases"]
    if len(cases) != 20 or any(any(v is not None for v in c["human_review"].values()) for c in cases):
        raise ValueError("Expected twenty unannotated development cases")

    wb = Workbook()
    review = wb.active
    review.title = "Revision"
    evidence = wb.create_sheet("Expedientes")
    for sheet in wb:
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    review.sheet_properties.tabColor = NAVY
    evidence.sheet_properties.tabColor = "8495A7"
    text_cell(review["A2"], "Revisión de 20 casos de desarrollo", bold=True, size=16)
    text_cell(review["A3"], "Abre el expediente y completa las celdas amarillas. No hay predicciones precargadas.", color=GRAY)
    text_cell(review["A4"], "Motivo: descríbelo en tus palabras. Evidencia: copia el fragmento que sustenta tu decisión.", color=GRAY)
    text_cell(review["A5"], "Causa: usa «No determinada» si no está sustentada. Marca revisión si el encaje o los datos son ambiguos.", color=GRAY)
    for row in (3, 4, 5):
        review.cell(row, 1).alignment = Alignment(wrap_text=False)
        review.row_dimensions[row].height = 21
    review["A2"].alignment = Alignment(wrap_text=False)
    review.row_dimensions[2].height = 28

    headers = ["Caso", "Radicado", "Expediente", "Tipología", "Motivo en tus palabras", "Evidencia textual",
               "Causa sustentada", "Datos o documentos pendientes", "Requiere revisión", "Observaciones"]
    widths = [9, 17, 18, 43, 43, 52, 42, 42, 18, 42]
    for col, (label, width) in enumerate(zip(headers, widths), 1):
        cell = review.cell(7, col, label)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        review.column_dimensions[cell.column_letter].width = width
    review.row_dimensions[7].height = 34
    review.freeze_panes = "D8"
    review.auto_filter.ref = "A7:J27"
    evidence.column_dimensions["A"].width = 32
    evidence.column_dimensions["B"].width = 124
    evidence.freeze_panes = "B7"
    text_cell(evidence["A2"], "Expedientes de desarrollo", bold=True, size=16)
    evidence["A2"].alignment = Alignment(wrap_text=False)
    text_cell(evidence["A3"], "Las notas se reproducen completas en fragmentos consecutivos. No se han resumido ni clasificado.", color=GRAY)
    evidence["A3"].alignment = Alignment(wrap_text=False)
    text_cell(evidence["A4"], "Fuentes: Informacion_Radicado_PTecnica.xlsx e Informacion_Cliente_PTecnica.xlsx", color=GRAY)
    evidence["A4"].alignment = Alignment(wrap_text=False)
    text_cell(evidence["A5"], "Se muestran todas las coincidencias de cliente. La presencia de una fila no prueba su asociación al caso.", color=GRAY)
    evidence["A5"].alignment = Alignment(wrap_text=False)
    for row in (3, 4, 5):
        evidence.row_dimensions[row].height = 23
    evidence.row_dimensions[2].height = 28

    row = 8
    note_cells = {}
    case_rows = {}
    for index, case in enumerate(cases):
        review_row = index + 8
        case_rows[case["case_id"]] = row
        source = case["radicado_source"]
        section(evidence, row, case["case_id"])
        evidence.cell(row, 2).value = f"Radicado {case['radicado']}"
        evidence.cell(row, 1).hyperlink = Hyperlink(ref=f"A{row}", location=f"'Revision'!A{review_row}")
        evidence.cell(row, 1).font = Font(name="Arial", size=11, color="FFFFFF", underline="single", bold=True)
        row += 1
        metadata = [("Fuente del radicado", f"Informacion_AQR!A{source['source_row']}:F{source['source_row']}")]
        metadata += [(key, source[key]) for key in payload["source_headers"]["radicados"] if key != "nota"]
        for key, value in metadata:
            text_cell(evidence.cell(row, 1), key, bold=True)
            text_cell(evidence.cell(row, 2), value)
            evidence.row_dimensions[row].height = 30
            row += 1
        note = source["nota"]
        parts = chunks(note)
        if "".join(parts) != note:
            raise AssertionError("Source note was altered")
        note_cells[case["case_id"]] = []
        for part_index, part in enumerate(parts, 1):
            text_cell(evidence.cell(row, 1), f"Nota {part_index}/{len(parts)}", bold=True)
            text_cell(evidence.cell(row, 2), part)
            evidence.cell(row, 2).fill = PatternFill("solid", fgColor="F5F7FA")
            evidence.row_dimensions[row].height = long_height(part)
            note_cells[case["case_id"]].append(f"B{row}")
            row += 1
        for candidate_index, candidate in enumerate(case["customer_candidates"], 1):
            section(evidence, row, f"Cliente {candidate_index}")
            evidence.cell(row, 2).value = f"Informacion_Cliente!A{candidate['source_row']}:H{candidate['source_row']}"
            row += 1
            for key in payload["source_headers"]["clientes"]:
                text_cell(evidence.cell(row, 1), key, bold=True)
                text_cell(evidence.cell(row, 2), candidate[key])
                evidence.row_dimensions[row].height = long_height(str(candidate[key]))
                row += 1
        row += 2

        text_cell(review.cell(review_row, 1), case["case_id"])
        text_cell(review.cell(review_row, 2), case["radicado"])
        text_cell(review.cell(review_row, 3), "Abrir expediente", color="0563C1")
        review.cell(review_row, 3).hyperlink = Hyperlink(
            ref=f"C{review_row}", location=f"'Expedientes'!A{case_rows[case['case_id']]}")
        review.cell(review_row, 3).font = Font(name="Arial", size=11, color="0563C1", underline="single")
        for col in range(4, 11):
            cell = review.cell(review_row, col)
            cell.fill = PatternFill("solid", fgColor=AMBER)
            cell.font = Font(name="Arial", size=11, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in range(1, 11):
            review.cell(review_row, col).border = Border(bottom=Side(style="hair", color="D5DCE5"))
        review.row_dimensions[review_row].height = 74

    # Catalog lives outside the annotation table. No model suggestion is selected.
    text_cell(review["A31"], "Tipologías permitidas", bold=True, size=13)
    review["A31"].alignment = Alignment(wrap_text=False)
    for i, label in enumerate(TIPOLOGIAS, 32):
        text_cell(review.cell(i, 1), label)
        review.cell(i, 1).alignment = Alignment(wrap_text=False)
        review.row_dimensions[i].height = 21
    dv = DataValidation(type="list", formula1="$A$32:$A$37", allow_blank=True)
    dv.errorTitle = "Elige una tipología del enunciado"
    dv.error = "Usa una de las seis opciones. Registra cualquier ambigüedad en las otras columnas."
    dv.showErrorMessage = True
    review.add_data_validation(dv)
    dv.add("D8:D27")
    flag = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    flag.showErrorMessage = True
    review.add_data_validation(flag)
    flag.add("I8:I27")

    wb.active = 0
    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)

    # Verify the saved artifact, including full-text retention and blank labels.
    check = load_workbook(args.output, data_only=False)
    for i, case in enumerate(cases, 8):
        if str(check["Revision"].cell(i, 2).value) != case["radicado"]:
            raise AssertionError("Incorrect review-to-case linkage")
        if any(check["Revision"].cell(i, c).value is not None for c in range(4, 11)):
            raise AssertionError("Human labels must remain blank")
        restored = "".join(check["Expedientes"][address].value or "" for address in note_cells[case["case_id"]])
        if restored != case["radicado_source"]["nota"]:
            raise AssertionError("Full source note not preserved in exported workbook")
    if len(check["Revision"].data_validations.dataValidation) != 2:
        raise AssertionError("Missing review dropdowns")
    if any(c.data_type == "f" for s in check for row_cells in s for c in row_cells):
        raise AssertionError("Review workbook should contain literal source text only")
    check.close()
    qa_dir = ROOT / ".work"
    qa_dir.mkdir(exist_ok=True)
    (qa_dir / "revision_layout.json").write_text(json.dumps({"case_rows": case_rows, "note_cells": note_cells}, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(args.output), "cases": len(cases), "sheets": wb.sheetnames,
                      "source_notes_preserved": True, "human_labels_blank": True}, ensure_ascii=False))


if __name__ == "__main__":
    main()
